import base64
from datetime import timedelta
from email import policy
from email.parser import BytesParser
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import Mock, patch
from zipfile import ZipFile

from django.test import TestCase, override_settings
from django.utils import timezone
from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from assistant.agent.documents import (
    DOCX_MIME, PDF_MIME, MAX_UPLOAD_BYTES, add_upload, email_me, extract_text,
    list_profile_documents, make_document, read_profile_document, save_profile_document,
    send_generated_document, send_telegram_documents,
)
from assistant.agent.loop import run_turn
from assistant.agent.proposals import ToolInputError
from assistant.agent.tools import TOOL_DEFINITIONS, ACTIVE_TOOL_DEFINITIONS, ToolContext
from assistant.agent.worker import claim_next_jobs, run_jobs
from assistant.services.gmail import GmailService
from assistant.services.telegram import TelegramAPIError, TelegramBot
from assistant.services.telegram_handlers import process_telegram_update
from assistant.tests.test_agent import AGENT_SETTINGS, AgentFakeBot, FakeClient, reply, text, tool_use
from assistant.tests import test_agent_email as email_tests
from assistant.tests.test_agent_email import FakeGmail, raw_email
from assistant.tests.test_telegram import message_update
from core.models import AgentEmailDelivery, AgentJob, AgentMessage, AgentThread, AgentUpload, GeneratedDocument, ProfileDocument


def docx_bytes(content="Consultant cardiologist"):
    document = Document()
    document.add_paragraph(content)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Experience"
    table.cell(0, 1).text = "NHS"
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def pdf_bytes():
    writer = PdfWriter()
    page = writer.add_blank_page(width=612, height=792)
    font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
    page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
    stream = DecodedStreamObject()
    stream.set_data(b"BT /F1 12 Tf 72 720 Td (Consultant cardiologist) Tj ET")
    page[NameObject("/Contents")] = writer._add_object(stream)
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


@override_settings(**AGENT_SETTINGS, AGENT_DOCUMENTS_ENABLED=True)
class DocumentAgentTests(TestCase):
    def setUp(self):
        self.bot = AgentFakeBot()
        self.bot.send_document = Mock(return_value={"message_id": 55})
        process_telegram_update(message_update(1, text="Help with my CV"), bot=self.bot)
        self.job = AgentJob.objects.get()
        self.thread = AgentThread.objects.create(channel="telegram", role="ike", telegram_user=self.job.telegram_user,
                                                 telegram_chat=self.job.telegram_chat, last_activity_at=timezone.now())
        self.job.thread = self.thread
        self.job.save()
        self.ctx = ToolContext(role="ike", speaker_name="Ike", thread=self.thread, job=self.job, bot=self.bot)
        self.gmail = FakeGmail()
        self.ctx.gmail_service = self.gmail

    def test_extracts_real_pdf_and_docx_including_tables(self):
        self.assertIn("Consultant cardiologist", extract_text(pdf_bytes(), PDF_MIME))
        extracted = extract_text(docx_bytes(), DOCX_MIME)
        self.assertIn("Consultant cardiologist", extracted)
        self.assertIn("Experience | NHS", extracted)

    def test_rejects_bad_oversized_and_expansion_bomb_files(self):
        for data, mime in [(b"invalid", PDF_MIME), (b"invalid", DOCX_MIME), (b"a" * (MAX_UPLOAD_BYTES + 1), PDF_MIME)]:
            with self.subTest(mime=mime), self.assertRaises(ToolInputError):
                extract_text(data, mime)
        stream = BytesIO()
        with ZipFile(stream, "w") as archive:
            for i in range(1001):
                archive.writestr(str(i), b"")
        with self.assertRaises(ToolInputError):
            extract_text(stream.getvalue(), DOCX_MIME)

    def upload(self, update_id, mime, data):
        payload = message_update(update_id, text="")
        payload["message"]["caption"] = "Save this as my CV"
        payload["message"]["document"] = {"file_id": "file", "file_name": "CV.pdf" if mime == PDF_MIME else "CV.docx", "mime_type": mime, "file_size": len(data)}
        self.bot.download_file = Mock(return_value=data)
        process_telegram_update(payload, bot=self.bot)
        return AgentJob.objects.get(telegram_update_id=update_id)

    def test_telegram_pdf_upload_queues_native_document_and_original_text(self):
        job = self.upload(2, PDF_MIME, pdf_bytes())
        self.assertEqual(job.input_blocks[-1]["type"], "document")
        self.assertEqual(base64.b64decode(job.input_blocks[-1]["source"]["data"]), pdf_bytes())
        self.assertIn("Consultant cardiologist", job.uploads.get().text)
        self.assertEqual(job.text, "Save this as my CV")
        self.bot.download_file.assert_called_once_with("voice/file_1.oga", max_bytes=MAX_UPLOAD_BYTES)

    def test_telegram_docx_upload_shows_extracted_text(self):
        job = self.upload(3, DOCX_MIME, docx_bytes())
        self.assertIn("Untrusted DOCX extracted text", job.input_blocks[-1]["text"])
        self.assertIn("Consultant cardiologist", job.input_blocks[-1]["text"])

    @override_settings(AGENT_DOCUMENTS_ENABLED=False)
    def test_upload_disabled_does_not_download(self):
        self.bot.get_file = Mock()
        payload = message_update(4, text="")
        payload["message"]["document"] = {"file_id": "file", "mime_type": PDF_MIME, "file_size": 20}
        process_telegram_update(payload, bot=self.bot)
        self.bot.get_file.assert_not_called()

    def test_oversized_telegram_metadata_rejected_before_download(self):
        payload = message_update(20, text="")
        payload["message"]["document"] = {"file_id": "file", "mime_type": PDF_MIME, "file_size": MAX_UPLOAD_BYTES + 1}
        self.bot.get_file = Mock()
        process_telegram_update(payload, bot=self.bot)
        self.bot.get_file.assert_not_called()
        self.assertIn("10 MB", self.bot.messages[-1][1])

    def test_group_documents_are_ignored(self):
        process_telegram_update(message_update(5, chat_id=-500, chat_type="group", text="/linkgroup"), bot=self.bot)
        payload = message_update(6, chat_id=-500, chat_type="group", text="")
        payload["message"]["document"] = {"file_id": "file", "mime_type": PDF_MIME, "file_size": 20}
        self.bot.get_file = Mock()
        process_telegram_update(payload, bot=self.bot)
        self.bot.get_file.assert_not_called()

    def test_profile_save_uses_latest_server_text_and_is_role_scoped(self):
        upload = add_upload(self.job, role="ike", filename="original.docx", mime_type=DOCX_MIME, data=docx_bytes())
        self.ctx.job = self.job
        result = save_profile_document(self.ctx, {"title": "My CV", "kind": "cv"})
        doc = ProfileDocument.objects.get(pk=result["profile_document_id"])
        self.assertEqual(doc.text, upload.text)
        self.assertEqual(doc.source_filename, "original.docx")
        self.assertEqual(list_profile_documents(self.ctx, {})["documents"][0]["id"], doc.id)
        self.assertEqual(read_profile_document(self.ctx, {"document_id": doc.id})["text"], upload.text)
        with self.assertRaises(ToolInputError):
            save_profile_document(self.ctx, {"title": "Fake", "kind": "cv", "text": "Invented"})
        other = ToolContext(role="wife", speaker_name="Wife", thread=self.thread, job=self.job)
        self.assertEqual(list_profile_documents(other, {})["documents"], [])
        with self.assertRaises(ToolInputError):
            read_profile_document(other, {"document_id": doc.id})
        with self.assertRaises(ToolInputError):
            save_profile_document(other, {"title": "Mine", "kind": "cv"})

    def test_generated_document_is_valid_docx_with_markup_and_can_be_resent(self):
        result = make_document(self.ctx, {"title": "Tailored CV", "content": "# Experience\n- **NHS** cardiology\n1. Teaching"})
        generated = GeneratedDocument.objects.get(pk=result["document_id"])
        opened = Document(BytesIO(bytes(generated.data)))
        self.assertEqual(opened.paragraphs[0].text, "Tailored CV")
        self.assertEqual(opened.paragraphs[1].style.name, "Heading 1")
        self.assertEqual(opened.paragraphs[2].style.name, "List Bullet")
        self.assertTrue(any(run.bold for run in opened.paragraphs[2].runs if run.text == "NHS"))
        send_telegram_documents(self.job, self.bot)
        send_telegram_documents(self.job, self.bot)
        self.bot.send_document.assert_called_once()
        send_generated_document(self.ctx, {"document_id": generated.id})
        self.assertEqual(self.job.document_deliveries.count(), 1)
        other = ToolContext(role="wife", speaker_name="Wife", thread=self.thread, job=self.job)
        with self.assertRaises(ToolInputError):
            send_generated_document(other, {"document_id": generated.id})

    def test_worker_delivers_document_after_text_through_streaming(self):
        # The worker selects an existing thread and current job from the queue.
        self.job.thread = None
        self.job.save()
        client = FakeClient(reply(tool_use("make_document", {"title": "CV", "content": "# Experience\nNHS"}), stop_reason="tool_use"), reply(text("Your CV is ready.")))
        run_jobs(claim_next_jobs(), bot=self.bot, client=client)
        self.assertEqual(self.bot.messages[-1][1], "Your CV is ready.")
        self.bot.send_document.assert_called_once()
        self.assertEqual(client.calls[0]["max_tokens"], 32000)
        self.assertEqual(client.calls[0]["thinking"], {"type": "adaptive", "block_binding": {"prefix_mismatch_behavior": "drop_block"}})
        for key in ["temperature", "top_p", "top_k", "budget_tokens"]:
            self.assertNotIn(key, client.calls[0])

    def test_email_me_recipient_lock_attachment_and_daily_limit(self):
        generated = make_document(self.ctx, {"title": "CV", "content": "NHS"})
        for i in range(10):
            email_me(self.ctx, {"subject": "My draft", "body": "Hello", "document_id": generated["document_id"]})
        self.assertEqual(len(self.gmail.sent), 10)
        self.assertTrue(all(message["to_addr"] == "ike@example.com" for message in self.gmail.sent))
        self.assertEqual(Document(BytesIO(self.gmail.sent[0]["attachments"][0]["data"])).paragraphs[0].text, "CV")
        with self.assertRaises(ToolInputError):
            email_me(self.ctx, {"subject": "Another", "body": "Hello"})
        with self.assertRaises(ToolInputError):
            email_me(self.ctx, {"subject": "A", "body": "B", "to_addr": "evil@example.com"})
        other = ToolContext(role="wife", speaker_name="Wife", gmail_service=self.gmail)
        email_me(other, {"subject": "Hers", "body": "Hello"})
        self.assertEqual(self.gmail.sent[-1]["to_addr"], "wife@example.com")
        with patch("assistant.agent.documents.timezone.localdate", return_value=timezone.localdate() + timedelta(days=1)):
            email_me(self.ctx, {"subject": "Tomorrow", "body": "Hello"})

    def test_uncertain_email_counts_against_limit(self):
        self.gmail.send_message = Mock(side_effect=RuntimeError("transport"))
        with self.assertRaises(ToolInputError):
            email_me(self.ctx, {"subject": "A", "body": "B"})
        self.assertEqual(AgentEmailDelivery.objects.count(), 1)

    @override_settings(AGENT_DOCUMENTS_ENABLED=False)
    def test_document_tools_fail_closed_when_disabled(self):
        with self.assertRaises(ToolInputError):
            make_document(self.ctx, {"title": "CV", "content": "Text"})
        with self.assertRaises(ToolInputError):
            email_me(self.ctx, {"subject": "A", "body": "B"})


@override_settings(**AGENT_SETTINGS, AGENT_EMAIL_ENABLED=True, AGENT_DOCUMENTS_ENABLED=True, AGENT_PROPOSALS_ENABLED=True)
class EmailDocumentTests(TestCase):
    poll = email_tests.EmailAgentTests.poll
    run_client = email_tests.EmailAgentTests.run_client

    def setUp(self):
        self.gmail = FakeGmail()

    def test_pdf_email_attachment_can_be_saved_from_server_text(self):
        msg = raw_email("Save this PDF as my CV")
        msg.add_attachment(pdf_bytes(), maintype="application", subtype="pdf", filename="profile.pdf")
        self.gmail.messages["pdfmail"] = (msg.as_bytes(), "thread-pdf")
        self.poll()
        client = FakeClient(reply(tool_use("save_profile_document", {"title": "PDF CV", "kind": "cv"}), stop_reason="tool_use"), reply(text("Saved.")))
        self.run_client(client)
        self.assertIn("Consultant cardiologist", ProfileDocument.objects.get().text)
        self.assertEqual(client.calls[0]["messages"][0]["content"][-1]["type"], "document")

    def test_email_attachment_can_be_saved_without_model_retyping_and_reply_has_docx(self):
        msg = raw_email("Save my CV and make me a cover letter")
        msg.add_attachment(docx_bytes(), maintype="application", subtype=DOCX_MIME.split("/")[1], filename="source.docx")
        self.gmail.messages["docmail"] = (msg.as_bytes(), "thread-docs")
        self.poll()
        client = FakeClient(reply(tool_use("save_profile_document", {"title": "Source CV", "kind": "cv"}, "a"),
                                  tool_use("make_document", {"title": "Cover letter", "content": "# Application\nI work in cardiology."}, "b"), stop_reason="tool_use"), reply(text("Your letter is attached.")))
        self.run_client(client)
        self.assertIn("Consultant cardiologist", ProfileDocument.objects.get().text)
        self.assertEqual(ProfileDocument.objects.get().source_filename, "source.docx")
        self.assertEqual(self.gmail.replies[0]["to_addr"], "ike@example.com")
        attachment = self.gmail.replies[0]["attachments"][0]
        self.assertEqual(Document(BytesIO(attachment["data"])).paragraphs[0].text, "Cover letter")
        self.assertTrue(AgentJob.objects.get().document_deliveries.get().sent_at)
        self.assertIn("Untrusted DOCX", str(client.calls[0]["messages"]))


class DocumentTransportTests(TestCase):
    @patch("assistant.services.gmail.get_google_credentials")
    @patch("assistant.services.gmail.build")
    def test_fresh_and_reply_mime_attachments(self, build, credentials):
        service = GmailService()
        attachment = {"filename": "cv.docx", "mime_type": DOCX_MIME, "data": docx_bytes()}
        for method in [service.send_message, service.send_reply]:
            kwargs = {"to_addr": "ike@example.com", "subject": "CV", "body_text": "Attached", "attachments": [attachment]}
            if method == service.send_reply:
                kwargs.update(thread_id="thread", in_reply_to_message_id_header="<m@example.com>")
            method(**kwargs)
            raw = build.return_value.users.return_value.messages.return_value.send.call_args.kwargs["body"]["raw"]
            msg = BytesParser(policy=policy.default).parsebytes(base64.urlsafe_b64decode(raw))
            part = list(msg.iter_attachments())[0]
            self.assertEqual(part.get_filename(), "cv.docx")
            self.assertEqual(Document(BytesIO(part.get_payload(decode=True))).paragraphs[0].text, "Consultant cardiologist")

    @patch("assistant.services.telegram.urlopen")
    def test_telegram_document_is_multipart_upload(self, urlopen):
        urlopen.return_value.__enter__.return_value.read.return_value = b'{"ok":true,"result":{"message_id":42}}'
        result = TelegramBot(token="private-token").send_document(123, filename="cv.docx", data=b"binary-docx")
        request = urlopen.call_args.args[0]
        self.assertTrue(request.full_url.endswith("/sendDocument"))
        self.assertIn("multipart/form-data", request.get_header("Content-type"))
        self.assertIn(b"binary-docx", request.data)
        self.assertIn(b'name="chat_id"', request.data)
        self.assertEqual(result["message_id"], 42)

    @patch("assistant.services.telegram.urlopen", side_effect=TimeoutError("secret URL private-token"))
    def test_telegram_document_errors_redact_token(self, urlopen):
        with self.assertRaises(TelegramAPIError) as caught:
            TelegramBot(token="private-token").send_document(123, filename="cv.docx", data=b"data")
        self.assertNotIn("private-token", str(caught.exception))


@override_settings(**AGENT_SETTINGS)
class WebStreamingTests(TestCase):
    def test_web_definitions_are_exact_sorted_and_no_code_execution(self):
        by_name = {t["name"]: t for t in TOOL_DEFINITIONS}
        self.assertEqual(by_name["web_search"], {"type": "web_search_20260209", "name": "web_search", "max_uses": 5})
        self.assertEqual(by_name["web_fetch"], {"type": "web_fetch_20260209", "name": "web_fetch", "max_uses": 5})
        self.assertNotIn("code_execution", by_name)
        self.assertEqual([t["name"] for t in TOOL_DEFINITIONS], sorted(by_name))

    def test_disabled_web_is_omitted_at_process_startup(self):
        self.assertNotIn("web_search", [t["name"] for t in ACTIVE_TOOL_DEFINITIONS])
        self.assertNotIn("web_fetch", [t["name"] for t in ACTIVE_TOOL_DEFINITIONS])

    def test_streaming_pause_server_error_and_history_are_replayed_verbatim(self):
        thread = AgentThread.objects.create(channel="email", role="ike", email_thread_id="web-thread", last_activity_at=timezone.now())
        AgentMessage.objects.create(thread=thread, role="user", content=[text("Read this advert")])
        AgentMessage.objects.create(thread=thread, role="system", content="context")
        server_reply = reply(
            {"type": "thinking", "thinking": "", "signature": "sig-web"},
            {"type": "server_tool_use", "id": "srv_1", "name": "web_search", "input": {"query": "public job advert"}},
            {"type": "web_search_tool_result", "tool_use_id": "srv_1", "content": {"type": "web_search_tool_result_error", "error_code": "unavailable"}},
            stop_reason="pause_turn",
        )
        client = FakeClient(server_reply, reply(text("Search is unavailable.")))
        with patch("assistant.agent.loop.ACTIVE_TOOL_DEFINITIONS", TOOL_DEFINITIONS):
            result = run_turn(thread, ToolContext(role="ike", speaker_name="Ike"), client=client)
        self.assertEqual(result.text, "Search is unavailable.")
        first, second = client.calls
        self.assertEqual(second["messages"][2]["content"], [b.to_dict(mode="json") for b in server_reply.content])
        self.assertEqual(first["system"], second["system"])
        self.assertEqual(first["tools"], second["tools"])
        self.assertEqual(first["cache_control"], {"type": "ephemeral", "ttl": "1h"})
        self.assertEqual(first["betas"], ["server-side-fallback-2026-07-01", "thinking-binding-controls-2026-08-01"])
        self.assertEqual(first["fallbacks"], "default")
